from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def export_timetable_xlsx(timetable_id: str, schedule: list[dict], output_dir: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Timetable"

    headers = ["Division", "Day", "Time Slot", "Subject", "Faculty", "Room", "Type"]
    ws.append(headers)
    header_fill = PatternFill(start_color="BDA6CE", end_color="BDA6CE", fill_type="solid")

    for idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in schedule:
        ws.append(
            [
                row.get("division"),
                row.get("day"),
                row.get("time_slot"),
                row.get("subject"),
                row.get("faculty"),
                row.get("room"),
                "Lab" if row.get("is_lab") else "Theory",
            ]
        )

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    file_path = str(Path(output_dir) / f"{timetable_id}.xlsx")
    wb.save(file_path)
    return file_path
